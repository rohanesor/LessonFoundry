"""Extract bounded text with stable locations; no OCR or legacy binary Office support."""

import io, zipfile
from pathlib import Path
from typing import Protocol
from fastapi import HTTPException


class DocumentExtractor(Protocol):
    def extract(self, name: str, data: bytes) -> list[tuple[str, str]]: ...


class FileExtractor:
    def extract(self, name, data):
        ext = Path(name).suffix.lower()
        if len(data) > 10 * 1024 * 1024:
            raise HTTPException(413, "Maximum file size is 10 MB")
        try:
            if ext in (".docx", ".pptx"):
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:
                        raise ValueError("Expanded Office file exceeds 50 MB")
            if ext == ".pdf":
                from pypdf import PdfReader

                doc = PdfReader(io.BytesIO(data))
                if len(doc.pages) > 150:
                    raise ValueError("Maximum 150 pages")
                rows = [
                    (f"Page {i + 1}", p.extract_text() or "")
                    for i, p in enumerate(doc.pages)
                ]
            elif ext == ".pptx":
                from pptx import Presentation

                rows = [
                    (
                        f"Slide {i + 1}",
                        "\n".join(s.text for s in slide.shapes if s.has_text_frame),
                    )
                    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides)
                ]
            elif ext == ".docx":
                from docx import Document

                doc = Document(io.BytesIO(data))
                rows = [
                    (f"Paragraph {i + 1}", p.text) for i, p in enumerate(doc.paragraphs)
                ]
                rows += [
                    (
                        f"Table {i + 1}",
                        "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows),
                    )
                    for i, t in enumerate(doc.tables)
                ]
            elif ext in (".txt", ".md"):
                rows = [("Text", data.decode("utf-8"))]
            else:
                raise ValueError(
                    "Supported formats: text PDF, PPTX, DOCX, TXT and Markdown"
                )
            if sum(len(t) for _, t in rows) > 150000:
                raise ValueError(
                    "Extracted text exceeds 150,000 characters; split the source"
                )
            if sum(len(t.strip()) for _, t in rows) < 30:
                raise ValueError(
                    "No usable text extracted. Scanned files require OCR; upload a text-based document."
                )
            return [(loc, t.strip()) for loc, t in rows if t.strip()]
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(422, f"Extraction failed: {str(e)[:240]}")
